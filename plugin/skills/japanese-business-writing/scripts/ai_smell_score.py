# -*- coding: utf-8 -*-
"""AI 臭さスコア — 資料の日本語を、SKILL.md の規則に当てて 0〜100 点で出す（低いほど人が書いた資料に近い）。

対象: .pptx（全テキスト枠）／.html（タグを落とした本文）／.md .txt（行）／.js（data.js の文字列リテラル）
使い方: python3 ai_smell_score.py 資料.pptx [他のファイル...] [--detail] [--speech]
  --detail 指摘の行を全部出す ／ --speech 発表原稿（敬体と否定終わりを数えない）

出し方:
  1,000 字あたりの重み付き件数 = Σ(重み × 件数) ÷ (総字数 ÷ 1000)
  スコア = min(100, round(重み付き件数 × 4))
  判定: 0〜9 納品可 ／ 10〜24 再点検（指摘の行を直す）／ 25〜 書き直し

重みは SKILL.md の章に対応する（数字は 2026-09 の登壇資料 41 枚の校正で決めた初期値）。
"""
import re, sys, os, collections

# 1) 語彙（C-2 口語、C-20 述語の和語逃げ、E-5 AI に多い語、依頼者の指定語）  重み 3
WORDS = ["そのまま", "かどうか", "まさに", "に至る", "という点で", "重要なのは", "いわば", "と言えるでしょう", "が求められます",
         "効いて", "効く", "はらむ", "出どころ", "やること", "要る", "高くつく", "手が届く", "溜まる", "潰す", "取られる",
         "ざっくり", "きちんと", "しっかり", "ちゃんと", "かなり", "ほぼほぼ", "やっていく", "見ていく", "持っていく",
         "肝", "鍵を握る", "カギ", "落とし穴", "転ばぬ先", "一石二鳥", "両輪", "武器", "追い風", "土俵"]
# 2) 型（A-17 理由の後置、E-3 ではなく、B-3 先取りの否定、C-18 敬体、A-6 つなぎ）  重み 2
PATTERNS = [
    (r"ためです|からです", "理由の後置（A-17）"),
    (r"ではなく", "AではなくB（E-3）"),
    (r"一方で|と言える", "上限語（H-1）"),
    (r"(ます|ました|ません|でした|です|ください)(。|$)", "スライド面の敬体（C-18）"),
    (r"^(まず|また|なお|さらに|そして|加えて|一方|つまり|ただ|しかし)、", "文頭のつなぎ（A-6）"),
    (r"ご安心|心配(は)?(あり|不要|無用)|懸念は(あり|ない)", "言われていない懸念の否定（B-3）"),
    (r"[をは](外し|足し|付け|変え)(た|、)|の箱[をに]", "変更報告の和語述語（C-22）"),
]
# 3) 誇張と抽象（A-7、A-4、E-8）  重み 2
HYPE = [r"大幅に|劇的に|飛躍的に|圧倒的|最大限|確実に|非常に|極めて|画期的|抜本的|徹底的|万全",
        r"[一-龥ァ-ン]+力(を|の|が)(向上|強化|底上げ)", r"(最適化|効率化|高度化|可視化|活性化)(を|に|し)"]
# 4) 構造（A-8 漢語の連続、A-10 連体修飾の重ね、F-4）  重み 1
MODIFY = re.compile(r"(による|における|に向けた|を通じ)")
KANGO_RUN = re.compile(r"[一-龥]{10,}")  # 漢字が 10 字以上続く（9 字の固有名詞は除く）
NEG_END = re.compile(r"(ない|ません|ず)$")
# 5) 表記（D 章。英数字と日本語の間の空白、全角英数字）  重み 0.5
SPACING = re.compile(r"(?<![・〜「」（）])[぀-ヿ一-鿿](?<![・〜「」（）])[0-9A-Za-z%]|[0-9A-Za-z%](?![・〜「」（）])[぀-ヿ一-鿿]")
ZENKAKU = re.compile(r"[０-９Ａ-Ｚａ-ｚ]")
SYMBOLS = re.compile(r"[！✓✔★☆■▶►◆●♪]|→.*→")

W_WORD, W_PAT, W_HYPE, W_STRUCT, W_NOTATION = 3, 2, 2, 1, 0.5
JP = re.compile(r"[぀-ヿ一-鿿]")


def lines_from(path):
    ext = os.path.splitext(path)[1].lower()
    if ext == ".pptx":
        from pptx import Presentation
        out = []
        for i, sl in enumerate(Presentation(path).slides, 1):
            for sp in sl.shapes:
                if sp.has_text_frame:
                    for p in sp.text_frame.paragraphs:
                        t = "".join(r.text for r in p.runs).strip()
                        if t: out.append(("p%d" % i, t))
                if getattr(sp, "has_table", False) and sp.has_table:
                    for row in sp.table.rows:
                        for c in row.cells:
                            t = c.text.strip()
                            if t: out.append(("p%d" % i, t))
        return out
    s = open(path, encoding="utf-8").read()
    if ext in (".html", ".htm"):
        s = re.sub(r"<(script|style)[^>]*>.*?</\1>", "", s, flags=re.S)
        s = re.sub(r"<br\s*/?>", "\n", s)
        s = re.sub(r"<[^>]+>", "\n", s)
        s = s.replace("&amp;", "&").replace("&lt;", "<").replace("&gt;", ">")
        return [("html", t.strip()) for t in s.split("\n") if JP.search(t)]
    if ext == ".js":
        lits = re.findall(r"'((?:[^'\\\n]|\\.)*)'", s)
        out = []
        for lit in lits:
            for t in lit.replace("\\n", "\n").split("\n"):
                if JP.search(t): out.append(("js", t.strip()))
        return out
    return [("L%d" % n, t.strip()) for n, t in enumerate(s.split("\n"), 1) if JP.search(t)]


def score(lines, detail=False, speech=False):
    chars = sum(len(t) for _, t in lines)
    hits = collections.OrderedDict()
    notes = []

    def add(cat, w, where, t, what):
        hits.setdefault(cat, [0, w]); hits[cat][0] += 1
        notes.append((cat, where, what, t[:50]))

    sent_total = sent_neg = 0
    for where, t in lines:
        for wd in WORDS:
            if wd in t: add("語彙", W_WORD, where, t, wd)
        for pat, name in PATTERNS:
            if speech and "敬体" in name: continue
            if re.search(pat, t): add("型", W_PAT, where, t, name)
        for pat in HYPE:
            for m in re.finditer(pat, t): add("誇張・抽象", W_HYPE, where, t, m.group())
        if len(MODIFY.findall(t)) >= 2: add("構造", W_STRUCT, where, t, "連体修飾 2 つ以上")
        for m in KANGO_RUN.finditer(t): add("構造", W_STRUCT, where, t, "漢語の連続 " + m.group())
        for sent in re.split(r"(?<=。)", t):
            s2 = sent.rstrip("。").strip()
            if not s2: continue
            sent_total += 1
            if NEG_END.search(s2): sent_neg += 1
        for m in SPACING.finditer(t): add("表記", W_NOTATION, where, t, "空白なし " + m.group())
        if ZENKAKU.search(t): add("表記", W_NOTATION, where, t, "全角英数字")
        if SYMBOLS.search(t): add("表記", W_NOTATION, where, t, "装飾記号")
    # 否定終わりは比率で見る（実資料 0.13%）。0.5% を超えた分を構造の件数に足す
    neg_rate = 100.0 * sent_neg / max(sent_total, 1)
    if neg_rate > 0.5 and not speech:
        extra = int((neg_rate - 0.5) * sent_total / 100)
        hits.setdefault("構造", [0, W_STRUCT]); hits["構造"][0] += extra
    weighted = sum(n * w for n, w in hits.values())
    per_k = weighted / max(chars / 1000.0, 0.001)
    sc = min(100, round(per_k * 4))
    verdict = "納品可" if sc < 10 else ("再点検" if sc < 25 else "書き直し")
    print("字数 %d ／ 行 %d ／ 否定終わり %.2f%%" % (chars, len(lines), neg_rate))
    for cat, (n, w) in hits.items():
        print("  %-6s %4d 件 × 重み %.1f = %5.1f" % (cat, n, w, n * w))
    print("重み付き件数 %.1f ／ 1,000 字あたり %.2f" % (weighted, per_k))
    print("AI 臭さスコア %d ／ 100 → %s" % (sc, verdict))
    if detail:
        print("--- 指摘の行 ---")
        for cat, where, what, t in notes:
            print("  [%s] %-5s %-28s %s" % (cat, where, what, t))
    return sc


if __name__ == "__main__":
    args = [a for a in sys.argv[1:] if not a.startswith("--")]
    detail = "--detail" in sys.argv
    speech = "--speech" in sys.argv  # 発表原稿（話し言葉の敬体と否定終わりを数えない）
    if not args:
        print(__doc__); sys.exit(1)
    for p in args:
        print("=== %s ===" % os.path.basename(p))
        score(lines_from(p), detail, speech)
