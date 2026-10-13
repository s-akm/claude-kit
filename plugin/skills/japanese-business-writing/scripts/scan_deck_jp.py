# -*- coding: utf-8 -*-
"""生成した pptx から日本語を全部抜き、実測の規則に当てる。

規則の出どころは ../ref/PHRASE.md（実企業スライド99枚の採集）と
../ref/FINDINGS.md（IR資料10社580ページの全文検索）。
"""
import re, sys, collections
from pptx import Presentation

NS = "{http://schemas.openxmlformats.org/drawingml/2006/main}"


def runs_of(prs):
    out = []
    for i, sl in enumerate(prs.slides, 1):
        for sp in sl.shapes:
            if not sp.has_text_frame:
                continue
            for p in sp.text_frame.paragraphs:
                txt = "".join(r.text for r in p.runs)
                if not txt.strip():
                    continue
                sz = max([(r.font.size.pt if r.font.size else 0)
                          for r in p.runs] or [0])
                faces = set()
                for r in p.runs:
                    rPr = r._r.find(NS + "rPr")
                    if rPr is None:
                        continue
                    for tag in ("latin", "ea", "cs"):
                        el = rPr.find(NS + tag)
                        if el is not None:
                            faces.add(el.get("typeface"))
                out.append((i, sz, txt, faces))
    return out


# 99枚で0件だった語（PHRASE.md）
ZERO = ["そのまま", "かどうか", "まさに", "に至る", "という点で",
        "重要なのは", "いわば", "と言えるでしょう", "が求められます"]
# 上限のある語（全文で数える。見出しだけでは足りない）
LIMITED = [("ではなく", 2), ("一方で", 1), ("と言える", 0), ("ためです", 0), ("からです", 0)]
KEITAI = re.compile(r"(ます|ました|ません|でした|です|ください)(。|$)")
# 図の中では使わない型
NEG_END = re.compile(r"(ない|ません|ず)$")
MODIFY = re.compile(r"(による|における|に向けた|を通じ)")


def w(ch):
    o = ord(ch)
    return 0.55 if (o < 0x2E80 or 0xFF61 <= o <= 0xFF9F) else 1.0


def width(s):
    return sum(w(c) for c in s)


def main(path):
    prs = Presentation(path)
    rs = runs_of(prs)
    print("=== 母数 ===")
    print("スライド %d 枚 ／ 文の行 %d 本" % (len(prs.slides.__iter__.__self__._sldIdLst), len(rs)))

    # ---- 1) フォント
    faces = collections.Counter()
    for _, _, _, f in rs:
        for x in f:
            faces[x] += 1
    print("\n=== フォント ===")
    for k, v in faces.most_common():
        print("  %-14s %d" % (k, v))

    # ---- 2) 級数
    print("\n=== 級数の分布 ===")
    sz = collections.Counter(round(s, 1) for _, s, _, _ in rs)
    for k in sorted(sz, reverse=True):
        print("  %5.1f pt  %3d本" % (k, sz[k]))
    small = [(i, s, t) for i, s, t, _ in rs if 0 < s < 8.0]
    print("  8pt未満: %d本" % len(small))
    for i, s, t in small[:6]:
        print("    p%-3d %.1fpt %s" % (i, s, t[:30]))

    # ---- 3) 使わないと決めた語
    print("\n=== 実資料99枚で0件だった語 ===")
    hit = 0
    for i, s, t, _ in rs:
        for z in ZERO:
            if z in t:
                print("  p%-3d %5.1fpt 「%s」 %s" % (i, s, z, t[:46]))
                hit += 1
    print("  計 %d件" % hit)

    # ---- 3-2) 上限のある語と、スライド面の敬体
    print("\n=== 上限のある語（資料全体で数える） ===")
    for z, lim in LIMITED:
        hits = [(i, t) for i, s, t, _ in rs if z in t]
        mark = "" if len(hits) <= lim else "  ← 上限 %d を超過" % lim
        print("  「%s」 %d件%s" % (z, len(hits), mark))
        for i, t in hits[:6]:
            print("      p%-3d %s" % (i, t[:46]))
    print("\n=== スライド面の敬体（0 が基準。原稿は別ファイル） ===")
    k = [(i, t) for i, s, t, _ in rs if KEITAI.search(t.strip())]
    for i, t in k:
        print("  p%-3d %s" % (i, t[:46]))
    print("  計 %d件" % len(k))

    # ---- 4) 否定終わり（文単位。実資料は4,720本中6本＝0.13%）
    print("\n=== 否定で終わる文（実資料 0.13%） ===")
    n = 0
    tot = 0
    for i, s, t, _ in rs:
        for sent in re.split(r"(?<=。)", t):
            s2 = sent.rstrip("。").strip()
            if not s2:
                continue
            tot += 1
            if NEG_END.search(s2):
                print("  p%-3d %5.1fpt %s" % (i, s, s2[:52]))
                n += 1
    print("  計 %d文 ／ 全%d文 = %.2f%%" % (n, tot, 100.0 * n / max(tot, 1)))

    # ---- 5) 連体修飾の重ね
    print("\n=== 連体修飾が2つ以上入った句 ===")
    n = 0
    for i, s, t, _ in rs:
        if len(MODIFY.findall(t)) >= 2:
            print("  p%-3d %s" % (i, t[:52]))
            n += 1
    print("  計 %d件" % n)

    # ---- 6) 字数（級数で要素を分ける）
    print("\n=== 字数（全角換算） ===")
    grp = collections.defaultdict(list)
    for i, s, t, _ in rs:
        if s >= 20: g = "見出し"
        elif s >= 13.5: g = "図の大見出し"
        elif s >= 12.0: g = "締めの帯・帯の中"
        elif s >= 10.0: g = "本文・説明"
        elif s > 0: g = "ラベル・脚注"
        else: continue
        grp[g].append((width(t), i, t))
    for g in ("見出し", "図の大見出し", "締めの帯・帯の中", "本文・説明",
              "ラベル・脚注"):
        v = sorted(grp[g])
        if not v:
            continue
        med = v[len(v) // 2][0]
        print("  %-12s %3d本  中央値 %4.1f字  最長 %4.1f字  p%d 「%s」"
              % (g, len(v), med, v[-1][0], v[-1][1], v[-1][2][:40]))

    # ---- 7) 数字の表記
    print("\n=== 数字の表記 ===")
    bad = []
    for i, s, t, _ in rs:
        for m in re.finditer(r"(?<![\d,.\-A-Za-z])\d{5,}(?![\d,.%年月日])", t):
            bad.append((i, m.group(), t[:40]))
    for i, g, t in bad:
        print("  p%-3d 桁区切りなし %s  %s" % (i, g, t))
    print("  計 %d件" % len(bad))


main(sys.argv[1] if len(sys.argv) > 1 else "AIDD_全編.pptx")
